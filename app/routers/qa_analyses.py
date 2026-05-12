"""Q&A Gap Analysis API — step 1 in Q&A → TVP → TC pipeline.

Endpoints:
- POST   /api/v1/projects/{pid}/qa-analyses/generate         (owner/pm/qc)
- GET    /api/v1/projects/{pid}/qa-analyses                   (any member)
- GET    /api/v1/projects/{pid}/qa-analyses/jobs/{job_id}/status
- GET    /api/v1/qa-analyses/{id}                             (any member)
- PATCH  /api/v1/qa-analyses/{id}/items/{item_id}             (any member)
- POST   /api/v1/qa-analyses/{id}/items                       (owner/pm/qc)
- DELETE /api/v1/qa-analyses/{id}/items/{item_id}             (owner/pm/qc)
"""

from __future__ import annotations

import time
import uuid
from typing import Annotated, Any

from fastapi import APIRouter, Depends, Path, status
from fastapi.responses import StreamingResponse
from redis.exceptions import RedisError
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from app.database import get_session
from app.deps import get_current_user
from app.exceptions import ApiError
from app.models.user import User
from app.project_access import is_system_admin, project_member_role, require_project_access
from app.redis_client import get_redis
from app.schemas.qa_analyses import (
    QAGapAnalysisDetailResponse,
    QAGapAnalysisListItem,
    QAGapAnalysisListResponse,
    QAGapAnalysisSummary,
    QAGapItemOut,
    QAGenerateAccepted,
    QAGenerateBody,
    QAItemCreateBody,
    QAItemDeleteResponse,
    QAItemPatchBody,
    QAJobProgressOut,
    QAJobResultOut,
    QAJobStatusResponse,
    QAUserBrief,
)
from app.worker import enqueue_qa_generate

router = APIRouter()


# --- Helpers --------------------------------------------------------------


async def _rate_limit_minute(redis, key: str, limit: int) -> None:
    bucket = int(time.time() // 60)
    rkey = f"{key}:{bucket}"
    try:
        n = await redis.incr(rkey)
        if n == 1:
            await redis.expire(rkey, 120)
    except RedisError as exc:
        raise ApiError(
            503,
            "REDIS_UNAVAILABLE",
            "Không kết nối được Redis. Kiểm tra REDIS_URL và dịch vụ Redis đang chạy.",
        ) from exc
    if n > limit:
        raise ApiError(429, "TOO_MANY_REQUESTS", "Quá nhiều yêu cầu. Vui lòng thử lại sau.")


def _can_generate_qa(user: User, role: str | None) -> bool:
    if is_system_admin(user):
        return True
    return (role or "").lower() in {"owner", "pm", "qc"}


async def _get_analysis_or_404(
    session: AsyncSession, qa_analysis_id: uuid.UUID
) -> dict[str, Any]:
    r = await session.execute(
        text(
            """
            SELECT id, project_id, screen_name, doc_types, status,
                   total_items, answered_items, generated_by, generated_at,
                   completed_at, created_at, updated_at
            FROM qa_gap_analyses WHERE id = :aid
            """
        ),
        {"aid": str(qa_analysis_id)},
    )
    row = r.mappings().first()
    if row is None:
        raise ApiError(404, "NOT_FOUND", "Q&A Gap Analysis không tồn tại")
    return dict(row)


async def _user_brief(session: AsyncSession, uid: uuid.UUID | None) -> QAUserBrief | None:
    if uid is None:
        return None
    r = await session.execute(
        text("SELECT id, full_name FROM users WHERE id = :uid"),
        {"uid": str(uid)},
    )
    row = r.mappings().first()
    if row is None:
        return None
    return QAUserBrief(id=uuid.UUID(str(row["id"])), full_name=str(row["full_name"] or ""))


async def _user_briefs(session: AsyncSession, uids: set[uuid.UUID]) -> dict[uuid.UUID, QAUserBrief]:
    if not uids:
        return {}
    rows = await session.execute(
        text("SELECT id, full_name FROM users WHERE id = ANY(CAST(:ids AS uuid[]))"),
        {"ids": [str(x) for x in uids]},
    )
    out: dict[uuid.UUID, QAUserBrief] = {}
    for row in rows.mappings():
        u = uuid.UUID(str(row["id"]))
        out[u] = QAUserBrief(id=u, full_name=str(row["full_name"] or ""))
    return out


def _progress_pct(total: int, answered: int) -> int:
    if total <= 0:
        return 0
    return int(round(100.0 * answered / total))


async def _recompute_analysis_counters(
    session: AsyncSession, analysis_id: uuid.UUID
) -> tuple[int, int, str]:
    """Update total_items, answered_items, status; return (total, answered, status)."""
    r = await session.execute(
        text(
            """
            SELECT
              COUNT(*) AS total,
              SUM(CASE WHEN answer_status = 'answered' THEN 1
                       WHEN answer_status IN ('wont_fix', 'deferred') THEN 1
                       ELSE 0 END) AS answered
            FROM qa_gap_items WHERE qa_analysis_id = CAST(:aid AS uuid)
            """
        ),
        {"aid": str(analysis_id)},
    )
    row = r.mappings().first() or {"total": 0, "answered": 0}
    total = int(row["total"] or 0)
    answered = int(row["answered"] or 0)
    new_status = "completed" if total > 0 and answered == total else "in_review" if answered > 0 else "draft"
    completed_at = "NOW()" if new_status == "completed" else "NULL"
    await session.execute(
        text(
            f"""
            UPDATE qa_gap_analyses
            SET total_items = :t,
                answered_items = :a,
                status = :st,
                completed_at = {completed_at},
                updated_at = NOW()
            WHERE id = CAST(:aid AS uuid)
            """
        ),
        {"t": total, "a": answered, "st": new_status, "aid": str(analysis_id)},
    )
    return total, answered, new_status


# --- Endpoints ------------------------------------------------------------


@router.post(
    "/projects/{project_id}/qa-analyses/generate",
    response_model=QAGenerateAccepted,
    status_code=status.HTTP_202_ACCEPTED,
)
async def enqueue_generate_qa(
    session: Annotated[AsyncSession, Depends(get_session)],
    user: Annotated[User, Depends(get_current_user)],
    project_id: Annotated[uuid.UUID, Path(...)],
    body: QAGenerateBody,
) -> QAGenerateAccepted:
    redis = get_redis()
    await _rate_limit_minute(redis, f"rl:qa:gen:{user.id}", 5)
    await require_project_access(session, user, project_id)
    role = await project_member_role(session, user.id, project_id)
    if not _can_generate_qa(user, role):
        raise ApiError(403, "FORBIDDEN", "Không có quyền generate Q&A")

    chk = await session.execute(
        text(
            """
            SELECT id FROM qa_jobs
            WHERE project_id = CAST(:pid AS uuid) AND status IN ('queued', 'processing')
            LIMIT 1
            """
        ),
        {"pid": str(project_id)},
    )
    if chk.scalar_one_or_none():
        raise ApiError(409, "CONFLICT", "Đang có job Q&A khác đang chạy")

    rchunks = await session.execute(
        text(
            """
            SELECT COUNT(*) FROM chunks c
            INNER JOIN doc_versions dv ON dv.id = c.doc_version_id AND dv.status = 'approved'
            INNER JOIN documents d ON d.id = dv.document_id
            WHERE d.project_id = CAST(:pid AS uuid)
              AND d.screen_name = :sn
              AND d.doc_type::text = ANY(:dts)
            """
        ),
        {"pid": str(project_id), "sn": body.screen_name, "dts": body.doc_types},
    )
    nchunks = int(rchunks.scalar_one() or 0)
    if nchunks == 0:
        raise ApiError(404, "NOT_FOUND", "Màn hình không có tài liệu approved cho doc_types đã chọn")

    if body.overwrite_existing:
        await session.execute(
            text(
                """
                DELETE FROM qa_gap_analyses
                WHERE project_id = CAST(:pid AS uuid) AND screen_name = :sn AND status = 'draft'
                """
            ),
            {"pid": str(project_id), "sn": body.screen_name},
        )

    job_id = uuid.uuid4()
    await session.execute(
        text(
            """
            INSERT INTO qa_jobs (
              id, project_id, screen_name, doc_types,
              status, progress, result, created_by, created_at, updated_at
            ) VALUES (
              CAST(:id AS uuid), CAST(:pid AS uuid), :sn, CAST(:dts AS text[]),
              'queued', '{}'::jsonb, '{}'::jsonb, CAST(:uid AS uuid), NOW(), NOW()
            )
            """
        ),
        {
            "id": str(job_id),
            "pid": str(project_id),
            "sn": body.screen_name,
            "dts": body.doc_types,
            "uid": str(user.id),
        },
    )
    await session.commit()

    enqueue_qa_generate(str(job_id))

    est = min(120, 20 + nchunks * 3)
    return QAGenerateAccepted(
        job_id=job_id,
        screen_name=body.screen_name,
        doc_types=body.doc_types,
        estimated_seconds=est,
        message="Đang phân tích Q&A. Vui lòng chờ...",
    )


@router.get(
    "/projects/{project_id}/qa-analyses/jobs/{job_id}/status",
    response_model=QAJobStatusResponse,
)
async def get_qa_job_status(
    session: Annotated[AsyncSession, Depends(get_session)],
    user: Annotated[User, Depends(get_current_user)],
    project_id: Annotated[uuid.UUID, Path(...)],
    job_id: Annotated[uuid.UUID, Path(...)],
) -> QAJobStatusResponse:
    redis = get_redis()
    await _rate_limit_minute(redis, f"rl:qa:jobst:{user.id}", 60)
    await require_project_access(session, user, project_id)

    jr = await session.execute(
        text(
            """
            SELECT id, project_id, status, progress, result, error_message,
                   qa_analysis_id, updated_at
            FROM qa_jobs WHERE id = CAST(:jid AS uuid)
            """
        ),
        {"jid": str(job_id)},
    )
    row = jr.mappings().first()
    if row is None or str(row["project_id"]) != str(project_id):
        raise ApiError(404, "NOT_FOUND", "Job không tồn tại")

    prog = row.get("progress") or {}
    if not isinstance(prog, dict):
        prog = {}
    total_chunks = int(prog.get("total_chunks") or 0)
    processed = int(prog.get("processed_chunks") or 0)
    pct = int(prog.get("percentage") or 0)

    res_out = None
    raw_res = row.get("result")
    if isinstance(raw_res, dict) and row.get("status") == "done":
        aid_raw = raw_res.get("qa_analysis_id") or row.get("qa_analysis_id")
        if aid_raw:
            try:
                res_out = QAJobResultOut(
                    qa_analysis_id=uuid.UUID(str(aid_raw)),
                    total_items=int(raw_res.get("total_items") or 0),
                )
            except (ValueError, TypeError):
                res_out = None

    st = str(row.get("status") or "queued")
    if st not in ("queued", "processing", "done", "failed"):
        st = "queued"

    return QAJobStatusResponse(
        job_id=job_id,
        status=st,
        progress=QAJobProgressOut(
            total_chunks=total_chunks,
            processed_chunks=processed,
            percentage=pct,
        ),
        result=res_out,
        error_message=str(row["error_message"]) if row.get("error_message") else None,
        updated_at=row.get("updated_at"),
    )


@router.get(
    "/projects/{project_id}/qa-analyses",
    response_model=QAGapAnalysisListResponse,
)
async def list_qa_analyses(
    session: Annotated[AsyncSession, Depends(get_session)],
    user: Annotated[User, Depends(get_current_user)],
    project_id: Annotated[uuid.UUID, Path(...)],
) -> QAGapAnalysisListResponse:
    redis = get_redis()
    await _rate_limit_minute(redis, f"rl:qa:list:{user.id}", 60)
    await require_project_access(session, user, project_id)

    rows = await session.execute(
        text(
            """
            SELECT id, project_id, screen_name, doc_types, status,
                   total_items, answered_items, generated_by, generated_at,
                   completed_at, created_at, updated_at
            FROM qa_gap_analyses
            WHERE project_id = CAST(:pid AS uuid)
            ORDER BY created_at DESC
            """
        ),
        {"pid": str(project_id)},
    )
    items = [dict(r) for r in rows.mappings().all()]

    user_ids: set[uuid.UUID] = set()
    for it in items:
        if it.get("generated_by"):
            user_ids.add(uuid.UUID(str(it["generated_by"])))
    briefs = await _user_briefs(session, user_ids)

    data: list[QAGapAnalysisListItem] = []
    for it in items:
        total = int(it.get("total_items") or 0)
        answered = int(it.get("answered_items") or 0)
        gb = it.get("generated_by")
        gen_user = briefs.get(uuid.UUID(str(gb))) if gb else None
        data.append(
            QAGapAnalysisListItem(
                id=uuid.UUID(str(it["id"])),
                project_id=uuid.UUID(str(it["project_id"])),
                screen_name=str(it.get("screen_name") or ""),
                doc_types=list(it.get("doc_types") or []),
                status=str(it.get("status") or "draft"),
                total_items=total,
                answered_items=answered,
                progress_percent=_progress_pct(total, answered),
                generated_by=gen_user,
                generated_at=it.get("generated_at"),
                completed_at=it.get("completed_at"),
                created_at=it.get("created_at"),
                updated_at=it.get("updated_at"),
            )
        )
    return QAGapAnalysisListResponse(data=data, total=len(data))


@router.get("/qa-analyses/{qa_analysis_id}", response_model=QAGapAnalysisDetailResponse)
async def get_qa_analysis_detail(
    session: Annotated[AsyncSession, Depends(get_session)],
    user: Annotated[User, Depends(get_current_user)],
    qa_analysis_id: Annotated[uuid.UUID, Path(...)],
) -> QAGapAnalysisDetailResponse:
    redis = get_redis()
    await _rate_limit_minute(redis, f"rl:qa:detail:{user.id}", 60)

    a = await _get_analysis_or_404(session, qa_analysis_id)
    project_id = uuid.UUID(str(a["project_id"]))
    await require_project_access(session, user, project_id)

    ir = await session.execute(
        text(
            """
            SELECT id, qa_analysis_id, gap_id, category, gap_description,
                   risk, question, answer, answer_status,
                   answered_by, answered_at, source_chunk_id, display_order,
                   created_at, updated_at
            FROM qa_gap_items
            WHERE qa_analysis_id = CAST(:aid AS uuid)
            ORDER BY display_order ASC, gap_id ASC
            """
        ),
        {"aid": str(qa_analysis_id)},
    )
    items_raw = [dict(r) for r in ir.mappings().all()]

    user_ids: set[uuid.UUID] = set()
    if a.get("generated_by"):
        user_ids.add(uuid.UUID(str(a["generated_by"])))
    for it in items_raw:
        if it.get("answered_by"):
            user_ids.add(uuid.UUID(str(it["answered_by"])))
    briefs = await _user_briefs(session, user_ids)

    items_out: list[QAGapItemOut] = []
    by_cat: dict[str, int] = {}
    by_risk: dict[str, int] = {}
    by_status: dict[str, int] = {}
    answered_count = 0
    for it in items_raw:
        cat = str(it.get("category") or "")
        rk = str(it.get("risk") or "")
        st = str(it.get("answer_status") or "open")
        by_cat[cat] = by_cat.get(cat, 0) + 1
        by_risk[rk] = by_risk.get(rk, 0) + 1
        by_status[st] = by_status.get(st, 0) + 1
        if st in ("answered", "wont_fix", "deferred"):
            answered_count += 1
        ab = it.get("answered_by")
        ab_user = briefs.get(uuid.UUID(str(ab))) if ab else None
        items_out.append(
            QAGapItemOut(
                id=uuid.UUID(str(it["id"])),
                qa_analysis_id=uuid.UUID(str(it["qa_analysis_id"])),
                gap_id=str(it.get("gap_id") or ""),
                category=cat,
                gap_description=str(it.get("gap_description") or ""),
                risk=rk,
                question=str(it.get("question") or ""),
                answer=str(it["answer"]) if it.get("answer") is not None else None,
                answer_status=st,
                answered_by=ab_user,
                answered_at=it.get("answered_at"),
                source_chunk_id=uuid.UUID(str(it["source_chunk_id"]))
                if it.get("source_chunk_id")
                else None,
                display_order=int(it.get("display_order") or 0),
                created_at=it.get("created_at"),
                updated_at=it.get("updated_at"),
            )
        )

    total = len(items_out)
    summary = QAGapAnalysisSummary(
        total_items=total,
        answered_items=answered_count,
        progress_percent=_progress_pct(total, answered_count),
        by_category=by_cat,
        by_risk=by_risk,
        by_status=by_status,
    )

    gen_brief = briefs.get(uuid.UUID(str(a["generated_by"]))) if a.get("generated_by") else None

    return QAGapAnalysisDetailResponse(
        id=uuid.UUID(str(a["id"])),
        project_id=project_id,
        screen_name=str(a.get("screen_name") or ""),
        doc_types=list(a.get("doc_types") or []),
        status=str(a.get("status") or "draft"),
        items=items_out,
        summary=summary,
        generated_by=gen_brief,
        generated_at=a.get("generated_at"),
        completed_at=a.get("completed_at"),
        created_at=a.get("created_at"),
        updated_at=a.get("updated_at"),
    )


@router.patch("/qa-analyses/{qa_analysis_id}/items/{item_id}", response_model=QAGapItemOut)
async def patch_qa_item(
    session: Annotated[AsyncSession, Depends(get_session)],
    user: Annotated[User, Depends(get_current_user)],
    qa_analysis_id: Annotated[uuid.UUID, Path(...)],
    item_id: Annotated[uuid.UUID, Path(...)],
    body: QAItemPatchBody,
) -> QAGapItemOut:
    redis = get_redis()
    await _rate_limit_minute(redis, f"rl:qa:patch:{user.id}", 120)

    a = await _get_analysis_or_404(session, qa_analysis_id)
    project_id = uuid.UUID(str(a["project_id"]))
    await require_project_access(session, user, project_id)

    chk = await session.execute(
        text(
            "SELECT id FROM qa_gap_items "
            "WHERE id = CAST(:iid AS uuid) AND qa_analysis_id = CAST(:aid AS uuid)"
        ),
        {"iid": str(item_id), "aid": str(qa_analysis_id)},
    )
    if chk.scalar_one_or_none() is None:
        raise ApiError(404, "NOT_FOUND", "Q&A item không tồn tại trong analysis này")

    if body.answer is None and body.answer_status is None:
        raise ApiError(400, "BAD_REQUEST", "Phải cung cấp answer hoặc answer_status")

    parts: list[str] = ["updated_at = NOW()"]
    params: dict[str, Any] = {"iid": str(item_id), "uid": str(user.id)}

    new_status: str | None = None
    if body.answer_status is not None:
        new_status = body.answer_status

    if body.answer is not None:
        parts.append("answer = :answer")
        params["answer"] = body.answer
        if new_status is None:
            new_status = "answered" if body.answer.strip() else "open"
        if new_status in ("answered", "wont_fix", "deferred"):
            parts.append("answered_by = CAST(:uid AS uuid)")
            parts.append("answered_at = NOW()")
        elif new_status == "open":
            parts.append("answered_by = NULL")
            parts.append("answered_at = NULL")

    if new_status is not None:
        parts.append("answer_status = :st")
        params["st"] = new_status
        if body.answer is None:
            if new_status in ("answered", "wont_fix", "deferred"):
                parts.append("answered_by = CAST(:uid AS uuid)")
                parts.append("answered_at = NOW()")
            elif new_status == "open":
                parts.append("answered_by = NULL")
                parts.append("answered_at = NULL")

    await session.execute(
        text(f"UPDATE qa_gap_items SET {', '.join(parts)} WHERE id = CAST(:iid AS uuid)"),
        params,
    )

    await _recompute_analysis_counters(session, qa_analysis_id)
    await session.commit()

    r = await session.execute(
        text(
            """
            SELECT id, qa_analysis_id, gap_id, category, gap_description,
                   risk, question, answer, answer_status,
                   answered_by, answered_at, source_chunk_id, display_order,
                   created_at, updated_at
            FROM qa_gap_items WHERE id = CAST(:iid AS uuid)
            """
        ),
        {"iid": str(item_id)},
    )
    it = r.mappings().first()
    if it is None:
        raise ApiError(500, "INTERNAL", "Không đọc lại được item sau khi update")

    ab_user = await _user_brief(session, uuid.UUID(str(it["answered_by"])) if it.get("answered_by") else None)

    return QAGapItemOut(
        id=uuid.UUID(str(it["id"])),
        qa_analysis_id=uuid.UUID(str(it["qa_analysis_id"])),
        gap_id=str(it.get("gap_id") or ""),
        category=str(it.get("category") or ""),
        gap_description=str(it.get("gap_description") or ""),
        risk=str(it.get("risk") or ""),
        question=str(it.get("question") or ""),
        answer=str(it["answer"]) if it.get("answer") is not None else None,
        answer_status=str(it.get("answer_status") or "open"),
        answered_by=ab_user,
        answered_at=it.get("answered_at"),
        source_chunk_id=uuid.UUID(str(it["source_chunk_id"]))
        if it.get("source_chunk_id")
        else None,
        display_order=int(it.get("display_order") or 0),
        created_at=it.get("created_at"),
        updated_at=it.get("updated_at"),
    )


@router.post(
    "/qa-analyses/{qa_analysis_id}/items",
    response_model=QAGapItemOut,
    status_code=status.HTTP_201_CREATED,
)
async def create_qa_item(
    session: Annotated[AsyncSession, Depends(get_session)],
    user: Annotated[User, Depends(get_current_user)],
    qa_analysis_id: Annotated[uuid.UUID, Path(...)],
    body: QAItemCreateBody,
) -> QAGapItemOut:
    redis = get_redis()
    await _rate_limit_minute(redis, f"rl:qa:add:{user.id}", 30)

    a = await _get_analysis_or_404(session, qa_analysis_id)
    project_id = uuid.UUID(str(a["project_id"]))
    await require_project_access(session, user, project_id)
    role = await project_member_role(session, user.id, project_id)
    if not _can_generate_qa(user, role):
        raise ApiError(403, "FORBIDDEN", "Không có quyền thêm gap")

    seq_r = await session.execute(
        text(
            "SELECT COALESCE(MAX(display_order), -1) + 1 AS n "
            "FROM qa_gap_items WHERE qa_analysis_id = CAST(:aid AS uuid)"
        ),
        {"aid": str(qa_analysis_id)},
    )
    next_order = int(seq_r.scalar_one() or 0)

    cnt_r = await session.execute(
        text(
            "SELECT COUNT(*) FROM qa_gap_items WHERE qa_analysis_id = CAST(:aid AS uuid)"
        ),
        {"aid": str(qa_analysis_id)},
    )
    next_gap_idx = int(cnt_r.scalar_one() or 0) + 1
    gap_id_str = f"G-{next_gap_idx:03d}"

    new_id = uuid.uuid4()
    now_answered = bool(body.answer and body.answer.strip())
    await session.execute(
        text(
            """
            INSERT INTO qa_gap_items (
              id, qa_analysis_id, gap_id, category, gap_description,
              risk, question, answer, answer_status,
              answered_by, answered_at, display_order,
              created_at, updated_at
            ) VALUES (
              CAST(:id AS uuid), CAST(:aid AS uuid), :gid, :cat, :desc,
              :risk, :q, :ans, :st,
              CASE WHEN :now_ans THEN CAST(:uid AS uuid) ELSE NULL END,
              CASE WHEN :now_ans THEN NOW() ELSE NULL END,
              :ord,
              NOW(), NOW()
            )
            """
        ),
        {
            "id": str(new_id),
            "aid": str(qa_analysis_id),
            "gid": gap_id_str,
            "cat": body.category,
            "desc": body.gap_description,
            "risk": body.risk,
            "q": body.question,
            "ans": body.answer,
            "st": "answered" if now_answered else "open",
            "now_ans": now_answered,
            "uid": str(user.id),
            "ord": next_order,
        },
    )

    await _recompute_analysis_counters(session, qa_analysis_id)
    await session.commit()

    r = await session.execute(
        text(
            """
            SELECT id, qa_analysis_id, gap_id, category, gap_description,
                   risk, question, answer, answer_status,
                   answered_by, answered_at, source_chunk_id, display_order,
                   created_at, updated_at
            FROM qa_gap_items WHERE id = CAST(:iid AS uuid)
            """
        ),
        {"iid": str(new_id)},
    )
    it = r.mappings().first()
    if it is None:
        raise ApiError(500, "INTERNAL", "Không đọc lại được item vừa tạo")

    ab_user = await _user_brief(session, uuid.UUID(str(it["answered_by"])) if it.get("answered_by") else None)

    return QAGapItemOut(
        id=uuid.UUID(str(it["id"])),
        qa_analysis_id=uuid.UUID(str(it["qa_analysis_id"])),
        gap_id=str(it.get("gap_id") or ""),
        category=str(it.get("category") or ""),
        gap_description=str(it.get("gap_description") or ""),
        risk=str(it.get("risk") or ""),
        question=str(it.get("question") or ""),
        answer=str(it["answer"]) if it.get("answer") is not None else None,
        answer_status=str(it.get("answer_status") or "open"),
        answered_by=ab_user,
        answered_at=it.get("answered_at"),
        source_chunk_id=uuid.UUID(str(it["source_chunk_id"]))
        if it.get("source_chunk_id")
        else None,
        display_order=int(it.get("display_order") or 0),
        created_at=it.get("created_at"),
        updated_at=it.get("updated_at"),
    )


@router.delete(
    "/qa-analyses/{qa_analysis_id}/items/{item_id}",
    response_model=QAItemDeleteResponse,
)
async def delete_qa_item(
    session: Annotated[AsyncSession, Depends(get_session)],
    user: Annotated[User, Depends(get_current_user)],
    qa_analysis_id: Annotated[uuid.UUID, Path(...)],
    item_id: Annotated[uuid.UUID, Path(...)],
) -> QAItemDeleteResponse:
    redis = get_redis()
    await _rate_limit_minute(redis, f"rl:qa:del:{user.id}", 30)

    a = await _get_analysis_or_404(session, qa_analysis_id)
    project_id = uuid.UUID(str(a["project_id"]))
    await require_project_access(session, user, project_id)
    role = await project_member_role(session, user.id, project_id)
    if not _can_generate_qa(user, role):
        raise ApiError(403, "FORBIDDEN", "Không có quyền xoá gap")

    res = await session.execute(
        text(
            "DELETE FROM qa_gap_items "
            "WHERE id = CAST(:iid AS uuid) AND qa_analysis_id = CAST(:aid AS uuid)"
        ),
        {"iid": str(item_id), "aid": str(qa_analysis_id)},
    )
    if res.rowcount == 0:
        raise ApiError(404, "NOT_FOUND", "Q&A item không tồn tại")

    await _recompute_analysis_counters(session, qa_analysis_id)
    await session.commit()

    return QAItemDeleteResponse(message="Đã xoá Q&A item", item_id=item_id)


@router.get("/qa-analyses/{qa_analysis_id}/export")
async def export_qa_analysis(
    session: Annotated[AsyncSession, Depends(get_session)],
    user: Annotated[User, Depends(get_current_user)],
    qa_analysis_id: Annotated[uuid.UUID, Path(...)],
) -> StreamingResponse:
    """Export Q&A items to Excel (.xlsx). Useful for sending to BA/Dev for review."""
    redis = get_redis()
    await _rate_limit_minute(redis, f"rl:qa:export:{user.id}", 10)

    a = await _get_analysis_or_404(session, qa_analysis_id)
    project_id = uuid.UUID(str(a["project_id"]))
    await require_project_access(session, user, project_id)

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as exc:
        raise ApiError(
            500,
            "MISSING_DEPENDENCY",
            "Thiếu thư viện openpyxl. Chạy `pip install openpyxl` trên server.",
        ) from exc

    ir = await session.execute(
        text(
            """
            SELECT gap_id, category, gap_description, risk, question, answer, answer_status,
                   display_order
            FROM qa_gap_items
            WHERE qa_analysis_id = CAST(:aid AS uuid)
            ORDER BY display_order ASC, gap_id ASC
            """
        ),
        {"aid": str(qa_analysis_id)},
    )
    items = [dict(r) for r in ir.mappings().all()]

    wb = Workbook()
    ws = wb.active
    if ws is None:
        raise ApiError(500, "INTERNAL", "Không tạo được worksheet")
    ws.title = "Q&A"

    headers = ["Gap ID", "Category", "Gap Description", "Risk", "Question", "Answer", "Status"]
    bold = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="6D28D9")
    wrap = Alignment(wrap_text=True, vertical="top")
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = bold
        cell.fill = fill
        cell.alignment = wrap

    widths = [12, 18, 60, 10, 60, 60, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w

    for row_idx, it in enumerate(items, start=2):
        ws.cell(row=row_idx, column=1, value=str(it.get("gap_id") or ""))
        ws.cell(row=row_idx, column=2, value=str(it.get("category") or ""))
        ws.cell(row=row_idx, column=3, value=str(it.get("gap_description") or ""))
        ws.cell(row=row_idx, column=4, value=str(it.get("risk") or ""))
        ws.cell(row=row_idx, column=5, value=str(it.get("question") or ""))
        ws.cell(row=row_idx, column=6, value=str(it.get("answer") or ""))
        ws.cell(row=row_idx, column=7, value=str(it.get("answer_status") or ""))
        for col in range(1, 8):
            ws.cell(row=row_idx, column=col).alignment = wrap

    ws.freeze_panes = "A2"

    import io

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    safe_screen = "".join(c if c.isalnum() or c in ("-", "_") else "_" for c in str(a.get("screen_name") or "screen"))
    filename = f"qa_{safe_screen}_{str(qa_analysis_id)[:8]}.xlsx"

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
